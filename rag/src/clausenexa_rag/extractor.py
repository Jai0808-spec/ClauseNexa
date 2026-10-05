from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Optional

from pypdf import PdfReader
from docx import Document


@dataclass
class ExtractionResult:
    text: str
    total_pages: Optional[int]


def extract_pdf(file_bytes: bytes) -> ExtractionResult:
    pdf_file = BytesIO(file_bytes)

    reader = PdfReader(pdf_file)

    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):
        text = page.extract_text()

        if text and text.strip():
            pages.append(
                f"[PAGE {page_number}]\n{text.strip()}"
            )

    return ExtractionResult(
        text="\n\n".join(pages),
        total_pages=len(reader.pages)
    )


def extract_docx(file_bytes: bytes) -> ExtractionResult:
    docx_file = BytesIO(file_bytes)

    document = Document(docx_file)

    content = []

    # Extract paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            content.append(text)

    # Extract tables
    for table in document.tables:

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                content.append(
                    " | ".join(cells)
                )

    # DOCX does not reliably expose page numbers
    # without rendering the document.
    return ExtractionResult(
        text="\n\n".join(content),
        total_pages=None
    )


def extract_document(
    file_bytes: bytes,
    file_name: str
) -> ExtractionResult:

    extension = Path(file_name).suffix.lower()

    if extension == ".pdf":
        return extract_pdf(file_bytes)

    if extension == ".docx":
        return extract_docx(file_bytes)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )